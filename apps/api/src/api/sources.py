"""??? / ?? / ?? API ???HIA-49 / HIA-55 / M1-02??

?? M1-02 ?????
- ??????``GET /evidences/project/{project_id}`` ????????
- ?????``Evidence.source_id`` + ``location`` + ``field_name`` + ``record_id``
  ?????? "???? ? ????" ???
- ?????????? ``require_role(Role.X)``?????/?? 404 ??
- ?????``align / reject / bulk-confirm / bulk-reject`` ???????
"""
from __future__ import annotations

import csv
import hashlib
import io
import uuid
from datetime import datetime, timezone
from typing import Optional

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Path,
    Query,
    Request,
    UploadFile,
    status,
)
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from src.db.connection import get_session
from src.db.evidence import (
    Evidence,
    EvidenceType,
    ProfilingRun,
    Source,
    SourceSnapshot,
    SourceStatus,
    SourceType,
)
from src.db.identity import Role
from src.db.governance import AuditEventType
from src.api.auth import (
    CurrentPrincipal,
    coerce_diff,
    record_audit,
    require_role,
    require_role_query,
)


# =====================================================================
# Source CRUD?????? + ???
# =====================================================================

router = APIRouter(prefix="/sources", tags=["data-sources"])


class SourceCreate(BaseModel):
    name: str = Field(..., min_length=1, max_length=255)
    description: Optional[str] = None
    source_type: SourceType
    connection_info: Optional[dict] = None
    access_scope: str = "restricted"


class SourceUpdate(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    status: Optional[SourceStatus] = None
    access_scope: Optional[str] = None


class SourceResponse(BaseModel):
    id: uuid.UUID
    project_id: uuid.UUID
    name: str
    description: Optional[str]
    source_type: SourceType
    status: SourceStatus
    file_size: Optional[int]
    row_count: Optional[int]
    column_count: Optional[int]
    access_scope: str
    created_at: str

    model_config = {"from_attributes": True}


class ProfilingRunResponse(BaseModel):
    id: uuid.UUID
    source_id: uuid.UUID
    status: str
    total_rows: Optional[int]
    sampled_rows: Optional[int]
    total_columns: Optional[int]
    error_message: Optional[str]
    started_at: Optional[str]
    completed_at: Optional[str]
    duration_ms: Optional[int]
    created_at: str

    model_config = {"from_attributes": True}


class EvidenceCreate(BaseModel):
    evidence_type: EvidenceType
    location: Optional[str] = None
    field_name: Optional[str] = None
    content: Optional[str] = None
    source_identifier: Optional[str] = None
    ontology_class_iri: Optional[str] = None
    property_iri: Optional[str] = None
    strength: str = "medium"
    notes: Optional[str] = None
    record_id: Optional[str] = None


class EvidenceResponse(BaseModel):
    id: uuid.UUID
    evidence_type: EvidenceType
    location: Optional[str]
    field_name: Optional[str]
    record_id: Optional[str]
    content: Optional[str]
    source_identifier: Optional[str]
    ontology_class_iri: Optional[str]
    property_iri: Optional[str]
    is_confirmed: bool
    strength: str
    created_at: str

    model_config = {"from_attributes": True}


# ---------- helpers ----------

async def _load_source_for_project(
    session: AsyncSession,
    *,
    source_id: uuid.UUID,
    project_id: uuid.UUID,
) -> Source:
    """????? source ???? project?????????? 404?"""
    result = await session.execute(
        select(Source).where(
            Source.id == source_id, Source.project_id == project_id
        )
    )
    s = result.scalar_one_or_none()
    if s is None:
        raise HTTPException(status_code=404, detail="source not found")
    return s


async def _load_evidence_for_project(
    session: AsyncSession,
    *,
    evidence_id: uuid.UUID,
    project_id: uuid.UUID,
) -> Evidence:
    result = await session.execute(
        select(Evidence).where(
            Evidence.id == evidence_id, Evidence.project_id == project_id
        )
    )
    e = result.scalar_one_or_none()
    if e is None:
        raise HTTPException(status_code=404, detail="evidence not found")
    return e


def _source_to_response(s: Source) -> SourceResponse:
    return SourceResponse(
        id=s.id,
        project_id=s.project_id,
        name=s.name,
        description=s.description,
        source_type=s.source_type,
        status=s.status,
        file_size=s.file_size,
        row_count=s.row_count,
        column_count=s.column_count,
        access_scope=s.access_scope,
        created_at=s.created_at.isoformat() if s.created_at else "",
    )


def _evidence_to_response(e: Evidence) -> EvidenceResponse:
    return EvidenceResponse(
        id=e.id,
        evidence_type=e.evidence_type,
        location=e.location,
        field_name=e.field_name,
        record_id=e.record_id,
        content=e.content,
        source_identifier=e.source_identifier,
        ontology_class_iri=e.ontology_class_iri,
        property_iri=e.property_iri,
        is_confirmed=e.is_confirmed,
        strength=e.strength,
        created_at=e.created_at.isoformat() if e.created_at else "",
    )


# ---------- Source ?? ----------

@router.get("", response_model=list[SourceResponse])
async def list_sources(
    project_id: uuid.UUID = Query(..., description="project scope"),
    source_type: Optional[SourceType] = Query(None),
    status_filter: Optional[SourceStatus] = Query(None, alias="status"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.VIEWER)),
    session: AsyncSession = Depends(get_session),
) -> list[SourceResponse]:
    """List sources scoped to a project. ``require_role`` ?? 404 ?????"""
    query = select(Source).where(Source.project_id == project_id)
    if source_type:
        query = query.where(Source.source_type == source_type)
    if status_filter:
        query = query.where(Source.status == status_filter)
    query = query.order_by(Source.created_at.desc())
    result = await session.execute(query)
    return [_source_to_response(s) for s in result.scalars().all()]


@router.post("", response_model=SourceResponse, status_code=status.HTTP_201_CREATED)
async def create_source(
    data: SourceCreate,
    request: Request,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.EDITOR)),
    session: AsyncSession = Depends(get_session),
) -> SourceResponse:
    """?????? Source?HIA-49 ??????????"""
    principal, _ = ctx
    source = Source(
        project_id=project_id,
        name=data.name,
        description=data.description,
        source_type=data.source_type,
        connection_info=data.connection_info,
        access_scope=data.access_scope,
        status=SourceStatus.UPLOADED,
        created_by=principal.user.id,
    )
    session.add(source)
    await session.flush()
    await session.refresh(source)
    await record_audit(
        session,
        event_type=AuditEventType.CREATE,
        principal=principal,
        project_id=project_id,
        request=request,
        target_type="source",
        target_id=str(source.id),
        target_label=source.name,
        after=coerce_diff(source),
    )
    return _source_to_response(source)


@router.get("/{source_id}", response_model=SourceResponse)
async def get_source(
    source_id: uuid.UUID,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.VIEWER)),
    session: AsyncSession = Depends(get_session),
) -> SourceResponse:
    """? source_id + project_id ????? Source?????"""
    s = await _load_source_for_project(
        session, source_id=source_id, project_id=project_id
    )
    return _source_to_response(s)


@router.patch("/{source_id}", response_model=SourceResponse)
async def update_source(
    source_id: uuid.UUID,
    data: SourceUpdate,
    request: Request,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.EDITOR)),
    session: AsyncSession = Depends(get_session),
) -> SourceResponse:
    principal, _ = ctx
    s = await _load_source_for_project(
        session, source_id=source_id, project_id=project_id
    )
    before = coerce_diff(s)
    for field, value in data.model_dump(exclude_unset=True).items():
        setattr(s, field, value)
    await session.flush()
    await session.refresh(s)
    await record_audit(
        session,
        event_type=AuditEventType.UPDATE,
        principal=principal,
        project_id=project_id,
        request=request,
        target_type="source",
        target_id=str(s.id),
        target_label=s.name,
        before=before,
        after=coerce_diff(s),
    )
    return _source_to_response(s)


@router.delete("/{source_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_source(
    source_id: uuid.UUID,
    request: Request,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.OWNER)),
    session: AsyncSession = Depends(get_session),
) -> None:
    principal, _ = ctx
    s = await _load_source_for_project(
        session, source_id=source_id, project_id=project_id
    )
    before = coerce_diff(s)
    label = s.name
    sid = str(s.id)
    await session.delete(s)
    await session.flush()
    await record_audit(
        session,
        event_type=AuditEventType.DELETE,
        principal=principal,
        project_id=project_id,
        request=request,
        target_type="source",
        target_id=sid,
        target_label=label,
        before=before,
    )


# ---------- ??????? source ??? ----------

@router.get("/{source_id}/evidences", response_model=list[EvidenceResponse])
async def list_evidences(
    source_id: uuid.UUID,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ontology_class_iri: Optional[str] = Query(None),
    property_iri: Optional[str] = Query(None),
    is_confirmed: Optional[bool] = Query(None),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.VIEWER)),
    session: AsyncSession = Depends(get_session),
) -> list[EvidenceResponse]:
    s = await _load_source_for_project(
        session, source_id=source_id, project_id=project_id
    )
    query = select(Evidence).where(
        Evidence.source_id == s.id, Evidence.project_id == project_id
    )
    if ontology_class_iri:
        query = query.where(Evidence.ontology_class_iri == ontology_class_iri)
    if property_iri:
        query = query.where(Evidence.property_iri == property_iri)
    if is_confirmed is not None:
        query = query.where(Evidence.is_confirmed == is_confirmed)
    query = query.order_by(Evidence.created_at.desc())
    result = await session.execute(query)
    return [_evidence_to_response(e) for e in result.scalars().all()]


@router.post(
    "/{source_id}/evidences", response_model=EvidenceResponse, status_code=status.HTTP_201_CREATED
)
async def create_evidence(
    source_id: uuid.UUID,
    data: EvidenceCreate,
    request: Request,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.EDITOR)),
    session: AsyncSession = Depends(get_session),
) -> EvidenceResponse:
    principal, _ = ctx
    s = await _load_source_for_project(
        session, source_id=source_id, project_id=project_id
    )
    evidence = Evidence(
        project_id=project_id,
        source_id=s.id,
        evidence_type=data.evidence_type,
        location=data.location,
        field_name=data.field_name,
        record_id=data.record_id,
        content=data.content,
        source_identifier=data.source_identifier,
        ontology_class_iri=data.ontology_class_iri,
        property_iri=data.property_iri,
        strength=data.strength,
        notes=data.notes,
        created_by=principal.user.id,
    )
    session.add(evidence)
    await session.flush()
    await session.refresh(evidence)
    await record_audit(
        session,
        event_type=AuditEventType.CREATE,
        principal=principal,
        project_id=project_id,
        request=request,
        target_type="evidence",
        target_id=str(evidence.id),
        target_label=data.field_name or data.content or "",
        after=coerce_diff(evidence),
    )
    return _evidence_to_response(evidence)


# ---------- ???? ----------

@router.get("/{source_id}/profiling", response_model=list[ProfilingRunResponse])
async def list_profiling_runs(
    source_id: uuid.UUID,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.VIEWER)),
    session: AsyncSession = Depends(get_session),
) -> list[ProfilingRunResponse]:
    s = await _load_source_for_project(
        session, source_id=source_id, project_id=project_id
    )
    result = await session.execute(
        select(ProfilingRun)
        .where(ProfilingRun.source_id == s.id)
        .order_by(ProfilingRun.created_at.desc())
    )
    return [
        ProfilingRunResponse(
            id=r.id,
            source_id=r.source_id,
            status=r.status.value if hasattr(r.status, "value") else str(r.status),
            total_rows=r.total_rows,
            sampled_rows=r.sampled_rows,
            total_columns=r.total_columns,
            error_message=r.error_message,
            started_at=r.started_at.isoformat() if r.started_at else None,
            completed_at=r.completed_at.isoformat() if r.completed_at else None,
            duration_ms=r.duration_ms,
            created_at=r.created_at.isoformat() if r.created_at else "",
        )
        for r in result.scalars().all()
    ]


@router.post(
    "/{source_id}/profiling", status_code=status.HTTP_201_CREATED
)
async def create_profiling_run(
    source_id: uuid.UUID,
    request: Request,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.EDITOR)),
    session: AsyncSession = Depends(get_session),
) -> dict:
    principal, _ = ctx
    s = await _load_source_for_project(
        session, source_id=source_id, project_id=project_id
    )
    profiling = ProfilingRun(
        source_id=s.id,
        status="pending",
        created_by=principal.user.id,
    )
    session.add(profiling)
    await session.flush()
    await session.refresh(profiling)
    await record_audit(
        session,
        event_type=AuditEventType.CREATE,
        principal=principal,
        project_id=project_id,
        request=request,
        target_type="profiling_run",
        target_id=str(profiling.id),
        target_label=f"source:{s.id}",
        after=coerce_diff(profiling),
    )
    return {
        "id": str(profiling.id),
        "status": profiling.status,
        "message": "profiling run created",
    }


# =====================================================================
# Evidence ?????? + ?????? prefix /evidences?
# =====================================================================


ev_router = APIRouter(prefix="/evidences", tags=["evidence"])


class EvidenceAlignInput(BaseModel):
    ontology_class_iri: str = Field(..., max_length=500)
    property_iri: Optional[str] = Field(None, max_length=500)
    confidence: Optional[float] = Field(None, ge=0.0, le=1.0)
    notes: Optional[str] = None


class EvidenceAlignResponse(BaseModel):
    id: uuid.UUID
    ontology_class_iri: str
    property_iri: Optional[str]
    alignment_confirmed: bool
    created_at: str


class EvidenceInboxResponse(BaseModel):
    id: uuid.UUID
    source_id: Optional[uuid.UUID]
    source_name: Optional[str]
    evidence_type: EvidenceType
    field_name: Optional[str]
    record_id: Optional[str]
    location: Optional[str]
    content: Optional[str]
    ontology_class_iri: Optional[str]
    property_iri: Optional[str]
    is_confirmed: bool
    strength: str
    created_at: str

    model_config = {"from_attributes": True}


class ProfilingExecuteRequest(BaseModel):
    sample_size: int = Field(default=1000, ge=100, le=50000)
    include_stats: bool = True
    detect_enums: bool = True


class FieldProfileResponse(BaseModel):
    field_name: str
    data_type: str
    total_count: int
    null_count: int
    null_ratio: float
    unique_count: int
    unique_ratio: float
    sample_values: list[str]
    min_value: Optional[str] = None
    max_value: Optional[str] = None
    avg_length: Optional[float] = None
    detected_enum_values: Optional[list[str]] = None
    suggested_ontology_type: Optional[str] = None
    suggested_property_type: Optional[str] = None


@ev_router.get("/project/{project_id}", response_model=list[EvidenceInboxResponse])
async def evidence_inbox(
    project_id: uuid.UUID = Path(..., description="project scope"),
    source_id: Optional[uuid.UUID] = Query(None),
    ontology_class_iri: Optional[str] = Query(None),
    is_confirmed: Optional[bool] = Query(None),
    search: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    size: int = Query(50, ge=1, le=200),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role(Role.VIEWER)),
    session: AsyncSession = Depends(get_session),
) -> list[EvidenceInboxResponse]:
    """????????????????????"""
    query = (
        select(Evidence, Source.name)
        .outerjoin(Source, Evidence.source_id == Source.id)
        .where(Evidence.project_id == project_id)
    )
    if source_id:
        query = query.where(Evidence.source_id == source_id)
    if ontology_class_iri:
        query = query.where(Evidence.ontology_class_iri == ontology_class_iri)
    if is_confirmed is not None:
        query = query.where(Evidence.is_confirmed == is_confirmed)
    if search:
        query = query.where(Evidence.content.ilike(f"%{search}%"))

    query = query.order_by(Evidence.created_at.desc())
    query = query.offset((page - 1) * size).limit(size)

    result = await session.execute(query)
    rows = result.all()

    return [
        EvidenceInboxResponse(
            id=e.id,
            source_id=e.source_id,
            source_name=source_name,
            evidence_type=e.evidence_type,
            field_name=e.field_name,
            record_id=e.record_id,
            location=e.location,
            content=e.content,
            ontology_class_iri=e.ontology_class_iri,
            property_iri=e.property_iri,
            is_confirmed=e.is_confirmed,
            strength=e.strength,
            created_at=e.created_at.isoformat() if e.created_at else "",
        )
        for e, source_name in rows
    ]


@ev_router.get("/{evidence_id}", response_model=EvidenceResponse)
async def get_evidence(
    evidence_id: uuid.UUID,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.VIEWER)),
    session: AsyncSession = Depends(get_session),
) -> EvidenceResponse:
    e = await _load_evidence_for_project(
        session, evidence_id=evidence_id, project_id=project_id
    )
    return _evidence_to_response(e)


@ev_router.patch("/{evidence_id}/align", response_model=EvidenceAlignResponse)
async def align_evidence(
    evidence_id: uuid.UUID,
    data: EvidenceAlignInput,
    request: Request,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.EDITOR)),
    session: AsyncSession = Depends(get_session),
) -> EvidenceAlignResponse:
    """?????????/???HIA-55 ??? ? confirm??

    - ?????????????
    - ???????
    """
    principal, _ = ctx
    e = await _load_evidence_for_project(
        session, evidence_id=evidence_id, project_id=project_id
    )
    before = coerce_diff(e)
    e.ontology_class_iri = data.ontology_class_iri
    e.property_iri = data.property_iri
    if data.confidence is not None:
        e.extraction_params = {
            **(e.extraction_params or {}),
            "confidence": data.confidence,
        }
    e.is_confirmed = True
    e.confirmed_by = principal.user.id
    e.confirmed_at = datetime.now(timezone.utc)
    if data.notes:
        e.notes = data.notes
    await session.flush()
    await session.refresh(e)
    await record_audit(
        session,
        event_type=AuditEventType.UPDATE,
        principal=principal,
        project_id=project_id,
        request=request,
        target_type="evidence",
        target_id=str(e.id),
        target_label=f"{e.field_name or e.content}@{data.ontology_class_iri}",
        before=before,
        after=coerce_diff(e),
    )
    return EvidenceAlignResponse(
        id=e.id,
        ontology_class_iri=e.ontology_class_iri,
        property_iri=e.property_iri,
        alignment_confirmed=e.is_confirmed,
        created_at=e.created_at.isoformat() if e.created_at else "",
    )


@ev_router.patch("/{evidence_id}/reject", response_model=EvidenceResponse)
async def reject_evidence(
    evidence_id: uuid.UUID,
    request: Request,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.EDITOR)),
    session: AsyncSession = Depends(get_session),
) -> EvidenceResponse:
    """?????HIA-55 ??? ? reject???????"""
    principal, _ = ctx
    e = await _load_evidence_for_project(
        session, evidence_id=evidence_id, project_id=project_id
    )
    before = coerce_diff(e)
    e.is_confirmed = False
    e.ontology_class_iri = None
    e.property_iri = None
    e.confirmed_by = principal.user.id
    e.confirmed_at = datetime.now(timezone.utc)
    await session.flush()
    await session.refresh(e)
    await record_audit(
        session,
        event_type=AuditEventType.UPDATE,
        principal=principal,
        project_id=project_id,
        request=request,
        target_type="evidence",
        target_id=str(e.id),
        target_label=e.field_name or e.content or "",
        before=before,
        after=coerce_diff(e),
    )
    return _evidence_to_response(e)


@ev_router.post("/bulk-confirm", status_code=status.HTTP_200_OK)
async def bulk_confirm_evidences(
    request: Request,
    payload: dict,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.EDITOR)),
    session: AsyncSession = Depends(get_session),
) -> dict:
    """?????????????? project ????? 1 ??????"""
    principal, _ = ctx
    evidence_ids_raw = payload.get("evidence_ids") or []
    ontology_class_iri: str = payload.get("ontology_class_iri", "")
    property_iri: Optional[str] = payload.get("property_iri")
    if not evidence_ids_raw or not ontology_class_iri:
        raise HTTPException(status_code=400, detail="evidence_ids and ontology_class_iri are required")
    try:
        evidence_ids = [uuid.UUID(str(x)) for x in evidence_ids_raw]
    except ValueError as e:
        raise HTTPException(status_code=400, detail=f"invalid evidence id: {e}")

    result = await session.execute(
        select(Evidence).where(
            Evidence.id.in_(evidence_ids), Evidence.project_id == project_id
        )
    )
    evidences = result.scalars().all()

    confirmed_at = datetime.now(timezone.utc)
    for e in evidences:
        e.ontology_class_iri = ontology_class_iri
        e.property_iri = property_iri
        e.is_confirmed = True
        e.confirmed_by = principal.user.id
        e.confirmed_at = confirmed_at

    await session.flush()
    await record_audit(
        session,
        event_type=AuditEventType.UPDATE,
        principal=principal,
        project_id=project_id,
        request=request,
        target_type="evidence_bulk",
        target_id=",".join(sorted(str(e.id) for e in evidences))[:255],
        target_label=f"bulk-confirm@{ontology_class_iri}",
        after={
            "count": len(evidences),
            "ontology_class_iri": ontology_class_iri,
            "property_iri": property_iri,
            "evidence_ids": [str(e.id) for e in evidences],
        },
        notes=f"requested={len(evidence_ids)} matched={len(evidences)}",
    )
    return {"updated": len(evidences), "total": len(evidence_ids)}


@ev_router.post("/bulk-reject", status_code=status.HTTP_200_OK)
async def bulk_reject_evidences(
    request: Request,
    payload: dict,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.EDITOR)),
    session: AsyncSession = Depends(get_session),
) -> dict:
    """???????? project ????? 1 ??????"""
    principal, _ = ctx
    evidence_ids_raw = payload.get("evidence_ids") or []
    if not evidence_ids_raw:
        raise HTTPException(status_code=400, detail="evidence_ids required")
    try:
        evidence_ids = [uuid.UUID(str(x)) for x in evidence_ids_raw]
    except ValueError as e:
        raise HTTPException(status_code=400, detail=f"invalid evidence id: {e}")

    result = await session.execute(
        select(Evidence).where(
            Evidence.id.in_(evidence_ids), Evidence.project_id == project_id
        )
    )
    evidences = result.scalars().all()

    confirmed_at = datetime.now(timezone.utc)
    for e in evidences:
        e.is_confirmed = False
        e.ontology_class_iri = None
        e.property_iri = None
        e.confirmed_by = principal.user.id
        e.confirmed_at = confirmed_at

    await session.flush()
    await record_audit(
        session,
        event_type=AuditEventType.UPDATE,
        principal=principal,
        project_id=project_id,
        request=request,
        target_type="evidence_bulk",
        target_id=",".join(sorted(str(e.id) for e in evidences))[:255],
        target_label="bulk-reject",
        after={
            "count": len(evidences),
            "evidence_ids": [str(e.id) for e in evidences],
        },
        notes=f"requested={len(evidence_ids)} matched={len(evidences)}",
    )
    return {"updated": len(evidences), "total": len(evidence_ids)}


# =====================================================================
# ???????HIA-49 ? CSV/XLSX/??/Markdown ?? + ?????
# =====================================================================


class SourceUploadResponse(BaseModel):
    source: SourceResponse
    snapshot_id: uuid.UUID
    row_count: int
    column_count: int
    auto_evidence_count: int
    auto_evidence_ids: list[uuid.UUID]


@router.post(
    "/upload",
    response_model=SourceUploadResponse,
    status_code=status.HTTP_201_CREATED,
)
async def upload_source(
    request: Request,
    file: UploadFile = File(...),
    project_id: uuid.UUID = Query(..., description="project scope"),
    description: Optional[str] = Query(None),
    auto_evidence: bool = Query(
        True,
        description="????????? SOURCE_FIELD ?????????????",
    ),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.EDITOR)),
    session: AsyncSession = Depends(get_session),
) -> SourceUploadResponse:
    """HIA-49 ? ?? CSV/XLSX/??/Markdown?????? Source + Snapshot + ?????

    - ?????? 32 MB???? 413?
    - ????? SourceSnapshot.checksum
    - ?????? ``SOURCE_FIELD`` ???location=``column:<name>``????????
    - ????
    """
    principal, _ = ctx

    MAX_BYTES = 32 * 1024 * 1024
    content = await file.read()
    if len(content) > MAX_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"file too large; max {MAX_BYTES} bytes",
        )
    if len(content) == 0:
        raise HTTPException(status_code=400, detail="empty file")

    filename = file.filename or "upload"
    source_type = _infer_source_type(filename, file.content_type)

    if source_type in (SourceType.CSV,):
        fields, row_count, sample_rows = _parse_csv(content)
    elif source_type in (SourceType.XLSX,):
        try:
            fields, row_count, sample_rows = _parse_xlsx(content)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"xlsx parse failed: {exc}")
    elif source_type in (SourceType.TEXT, SourceType.MARKDOWN):
        fields, row_count, sample_rows = _parse_text(content)
    elif source_type == SourceType.JSON:
        fields, row_count, sample_rows = _parse_json(content)
    else:
        fields, row_count, sample_rows = [], 0, []

    schema_info = {
        "filename": filename,
        "content_type": file.content_type,
        "fields": fields,
        "sample_rows": sample_rows,
    }

    checksum = hashlib.sha256(content).hexdigest()

    source = Source(
        project_id=project_id,
        name=filename,
        description=description,
        source_type=source_type,
        connection_info={"filename": filename, "content_type": file.content_type},
        file_path=None,
        file_size=len(content),
        row_count=row_count,
        column_count=len(fields),
        access_scope="restricted",
        is_sensitive=False,
        schema_info=schema_info,
        status=SourceStatus.PARSED if fields else SourceStatus.UPLOADED,
        created_by=principal.user.id,
    )
    session.add(source)
    await session.flush()
    await session.refresh(source)

    snapshot = SourceSnapshot(
        source_id=source.id,
        version=1,
        snapshot_type="upload",
        storage_path=None,
        storage_size=len(content),
        checksum=checksum,
        row_count=row_count,
        schema_hash=hashlib.sha256(
            (",".join(f.get("name", "") for f in fields)).encode("utf-8")
        ).hexdigest(),
        created_by=principal.user.id,
    )
    session.add(snapshot)
    await session.flush()
    await session.refresh(snapshot)

    auto_evidence_ids: list[uuid.UUID] = []
    if auto_evidence and fields:
        for f in fields:
            ev = Evidence(
                project_id=project_id,
                source_id=source.id,
                evidence_type=EvidenceType.SOURCE_FIELD,
                location=f"column:{f.get('name', '')}",
                field_name=f.get("name"),
                content=", ".join(f.get("sample_values", [])[:5]) or None,
                source_identifier=filename,
                extraction_method="upload_parse",
                extraction_params={
                    "inferred_type": f.get("type"),
                    "row": None,
                },
                strength="medium",
                created_by=principal.user.id,
            )
            session.add(ev)
        await session.flush()
        result = await session.execute(
            select(Evidence).where(
                Evidence.source_id == source.id,
                Evidence.project_id == project_id,
                Evidence.evidence_type == EvidenceType.SOURCE_FIELD,
            )
        )
        auto_evidence_ids = [e.id for e in result.scalars().all()]

    await record_audit(
        session,
        event_type=AuditEventType.CREATE,
        principal=principal,
        project_id=project_id,
        request=request,
        target_type="source",
        target_id=str(source.id),
        target_label=source.name,
        after={
            "filename": filename,
            "source_type": source_type.value,
            "row_count": row_count,
            "column_count": len(fields),
            "snapshot_id": str(snapshot.id),
            "checksum": checksum,
            "auto_evidence_count": len(auto_evidence_ids),
        },
        notes=f"size={len(content)} bytes",
    )

    return SourceUploadResponse(
        source=_source_to_response(source),
        snapshot_id=snapshot.id,
        row_count=row_count,
        column_count=len(fields),
        auto_evidence_count=len(auto_evidence_ids),
        auto_evidence_ids=auto_evidence_ids,
    )


def _infer_source_type(filename: str, content_type: Optional[str]) -> SourceType:
    name = filename.lower()
    if name.endswith(".csv"):
        return SourceType.CSV
    if name.endswith((".xlsx", ".xlsm")):
        return SourceType.XLSX
    if name.endswith(".json") or (content_type or "").startswith("application/json"):
        return SourceType.JSON
    if name.endswith((".md", ".markdown")):
        return SourceType.MARKDOWN
    if name.endswith((".ttl", ".owl", ".nt", ".rdf")):
        return SourceType.RDF
    if content_type and content_type.startswith("text/"):
        return SourceType.TEXT
    return SourceType.TEXT


def _parse_csv(content: bytes) -> tuple[list[dict], int, list[dict]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("gbk", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    rows: list[dict] = []
    for i, row in enumerate(reader):
        if i >= 100:
            break
        rows.append({k: (v if v is not None else "") for k, v in row.items()})
    fields = _infer_fields(rows)
    try:
        total = sum(1 for _ in csv.DictReader(io.StringIO(text)))
    except Exception:
        total = len(rows)
    return fields, total, rows


def _parse_xlsx(content: bytes) -> tuple[list[dict], int, list[dict]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl not installed; install via `pip install openpyxl`"
        ) from exc
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return [], 0, []
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = list(next(rows_iter))
    except StopIteration:
        return [], 0, []
    header = [str(c) if c is not None else "" for c in header]
    rows: list[dict] = []
    for i, r in enumerate(rows_iter):
        if i >= 100:
            break
        rows.append({header[j]: ("" if r[j] is None else str(r[j])) for j in range(len(header))})
    fields = _infer_fields(rows)
    total = ws.max_row or 0
    return fields, total, rows


def _parse_text(content: bytes) -> tuple[list[dict], int, list[dict]]:
    text = content.decode("utf-8", errors="replace")
    lines = text.splitlines()
    return (
        [{"name": "line", "type": "string", "sample_values": lines[:5]}],
        len(lines),
        [],
    )


def _parse_json(content: bytes) -> tuple[list[dict], int, list[dict]]:
    import json as _json
    text = content.decode("utf-8", errors="replace")
    try:
        data = _json.loads(text)
    except _json.JSONDecodeError as exc:
        raise RuntimeError(f"invalid json: {exc}") from exc
    if isinstance(data, list):
        rows = data[:100]
        total = len(data)
    elif isinstance(data, dict):
        rows = [data]
        total = 1
    else:
        rows = [{"value": str(data)}]
        total = 1
    fields = _infer_fields(rows)
    return fields, total, rows


def _infer_fields(rows: list[dict]) -> list[dict]:
    if not rows:
        return []
    field_names: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                field_names.append(k)
    out: list[dict] = []
    for name in field_names:
        values = [r.get(name) for r in rows if r.get(name) not in (None, "")]
        sample_values = [str(v)[:200] for v in values[:5]]
        inferred_type = _infer_scalar_type(values)
        out.append(
            {
                "name": name,
                "type": inferred_type,
                "sample_values": sample_values,
                "unique_count": len({str(v) for v in values}),
                "null_count": sum(1 for r in rows if r.get(name) in (None, "")),
            }
        )
    return out


def _infer_scalar_type(values: list) -> str:
    if not values:
        return "string"
    int_like = all(_is_int(v) for v in values)
    if int_like:
        return "integer"
    float_like = all(_is_float(v) for v in values)
    if float_like:
        return "float"
    bool_like = all(_is_bool(v) for v in values)
    if bool_like:
        return "boolean"
    date_like = all(_is_date(v) for v in values)
    if date_like:
        return "date"
    return "string"


def _is_int(v) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, int):
        return True
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return False
        try:
            int(s)
            return True
        except ValueError:
            return False
    return False


def _is_float(v) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return False
        try:
            float(s)
            return True
        except ValueError:
            return False
    return False


def _is_bool(v) -> bool:
    if isinstance(v, bool):
        return True
    if isinstance(v, str):
        return v.strip().lower() in ("true", "false", "yes", "no", "0", "1")
    return False


def _is_date(v) -> bool:
    if not isinstance(v, str):
        return False
    s = v.strip()
    if not s:
        return False
    for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            datetime.strptime(s, fmt)
            return True
        except ValueError:
            continue
    return False


# =====================================================================
# ????????? prefix /profiling?
# =====================================================================


prof_router = APIRouter(prefix="/profiling", tags=["profiling"])


@prof_router.post("/sources/{source_id}/execute", response_model=dict)
async def execute_field_profiling(
    source_id: uuid.UUID,
    data: ProfilingExecuteRequest,
    project_id: uuid.UUID = Query(..., description="project scope"),
    ctx: tuple[CurrentPrincipal, Role] = Depends(require_role_query(Role.EDITOR)),
    session: AsyncSession = Depends(get_session),
) -> dict:
    """????????? Source.schema_info ???? schema ?????"""
    s = await _load_source_for_project(
        session, source_id=source_id, project_id=project_id
    )
    schema = s.schema_info or {}
    fields = schema.get("fields", [])

    if not fields:
        conn = s.connection_info or {}
        if s.source_type == "postgresql" and "query" in conn:
            return {
                "source_id": str(source_id),
                "status": "pending_implementation",
                "message": "PostgreSQL profiling requires connector plugin",
                "fields": [],
            }
        return {
            "source_id": str(source_id),
            "status": "no_schema",
            "message": "source has no schema_info, please upload first",
            "fields": [],
        }

    profiles: list[dict] = []
    for field in fields[:50]:
        f_name = field.get("name", "")
        f_type = field.get("type", "string").lower()
        profile = {
            "field_name": f_name,
            "data_type": f_type,
            "total_count": data.sample_size,
            "null_count": field.get("null_count", 0),
            "null_ratio": field.get("null_ratio", 0.0),
            "unique_count": field.get("unique_count", 0),
            "unique_ratio": field.get("unique_ratio", 0.0),
            "sample_values": field.get("sample_values", [])[:10],
            "min_value": field.get("min_value"),
            "max_value": field.get("max_value"),
            "avg_length": field.get("avg_length"),
            "detected_enum_values": None,
            "suggested_ontology_type": _infer_ontology_type(f_type, field),
            "suggested_property_type": _infer_property_type(f_type),
        }
        if data.detect_enums:
            enum_vals = field.get("enum_values", [])
            if enum_vals:
                profile["detected_enum_values"] = enum_vals[:50]
            elif profile["unique_ratio"] < 0.05 and profile["unique_count"] <= 50:
                profile["detected_enum_values"] = field.get("sample_values", [])[:50]
        profiles.append(profile)

    run_result = await session.execute(
        select(ProfilingRun)
        .where(ProfilingRun.source_id == s.id)
        .order_by(ProfilingRun.created_at.desc())
        .limit(1)
    )
    run = run_result.scalar_one_or_none()
    if run:
        run.status = "completed"
        run.results = {"fields": profiles}
        run.total_rows = data.sample_size
        run.total_columns = len(profiles)
        await session.flush()

    return {
        "source_id": str(source_id),
        "run_id": str(run.id) if run else None,
        "status": "completed",
        "field_count": len(profiles),
        "fields": profiles,
    }


def _infer_ontology_type(field_type: str, field: dict) -> str:
    t = field_type.lower()
    if t in ("integer", "bigint", "smallint"):
        return "xsd:integer"
    if t in ("numeric", "decimal", "float", "double"):
        return "xsd:decimal"
    if t == "boolean":
        return "xsd:boolean"
    if t in ("date", "datetime", "timestamp"):
        return "xsd:dateTime"
    return "xsd:string"


def _infer_property_type(field_type: str) -> str:
    t = field_type.lower()
    if t in (
        "integer", "bigint", "smallint", "numeric", "decimal", "float", "double",
        "boolean", "date", "datetime", "timestamp", "json", "text", "string",
    ):
        return "datatype"
    return "datatype"
